import os
import glob
from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(HERE, "slides_png")
OUT = os.path.join(HERE, "Data Storytelling y diseno de dashboards.pptx")

images = sorted(glob.glob(os.path.join(IMG_DIR, "slide_*.png")))
print(f"Found {len(images)} images")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

for img_path in images:
    slide = prs.slides.add_slide(blank_layout)
    slide.shapes.add_picture(img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

prs.save(OUT)
print("Saved:", OUT)
